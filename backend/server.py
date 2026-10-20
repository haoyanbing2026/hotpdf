
import asyncio
import json
import logging
import sqlite3
import secrets
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
import os
import re
import shutil
import subprocess
import tempfile
import time
import uuid
import zipfile
import base64
from email.message import EmailMessage
from urllib import request as urlrequest, parse as urlparse, error as urlerror

from contextlib import asynccontextmanager, suppress, contextmanager
from pathlib import Path

import fitz

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile, Header, Request
from pydantic import BaseModel, Field, EmailStr, ValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from starlette.concurrency import run_in_threadpool


# ========================================
# Configuration
# ========================================

logger = logging.getLogger("hotpdf")

UPLOAD_DIR = Path(
    os.getenv(
        "UPLOAD_DIR",
        os.path.join(tempfile.gettempdir(), "hotpdf_files")
    )
).resolve()

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True,
    mode=0o700
)

MAX_FILE_MB = int(
    os.getenv("MAX_FILE_MB", "20")
)

MAX_MERGE_TOTAL_MB = int(
    os.getenv("MAX_MERGE_TOTAL_MB", "60")
)

MAX_MERGE_FILES = int(
    os.getenv("MAX_MERGE_FILES", "10")
)

MAX_PDF_PAGES = int(
    os.getenv("MAX_PDF_PAGES", "300")
)

MAX_MERGE_PAGES = int(
    os.getenv("MAX_MERGE_PAGES", "600")
)

TEMP_TTL_MINUTES = int(
    os.getenv("TEMP_TTL_MINUTES", "60")
)

CLEAN_INTERVAL_SECONDS = 600

MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024

MAX_MERGE_TOTAL_BYTES = (
    MAX_MERGE_TOTAL_MB * 1024 * 1024
)

TEMP_TTL_SECONDS = TEMP_TTL_MINUTES * 60


# Only server-generated filenames are accepted.

FILE_NAME_RE = re.compile(
    r"(?:upload|split|merged|compressed|converted|edited)_"
    r"[0-9a-f]{32}\.(?:pdf|docx|pptx|zip)"
)

PDF_NAME_RE = re.compile(
    r"(?:upload|split|merged|compressed|converted|edited)_"
    r"[0-9a-f]{32}\.pdf"
)


# ========================================
# Author feedback storage
# ========================================

# Use PostgreSQL on Render when DATABASE_URL is set. Keep SQLite locally.
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
USE_POSTGRES = bool(DATABASE_URL)
ADMIN_TOKEN = os.getenv("FEEDBACK_ADMIN_TOKEN", "").strip()
GMAIL_CLIENT_ID = os.getenv("GMAIL_CLIENT_ID", "").strip()
GMAIL_CLIENT_SECRET = os.getenv("GMAIL_CLIENT_SECRET", "").strip()
GMAIL_REFRESH_TOKEN = os.getenv("GMAIL_REFRESH_TOKEN", "").strip()
GMAIL_SENDER = os.getenv("GMAIL_SENDER", "").strip()

def gmail_api_configured():
    return all((GMAIL_CLIENT_ID, GMAIL_CLIENT_SECRET, GMAIL_REFRESH_TOKEN, GMAIL_SENDER))

if not USE_POSTGRES:
    FEEDBACK_DB = Path(os.getenv("FEEDBACK_DB_PATH", "/data/hotpdf_feedback.sqlite3"))
    FEEDBACK_DB.parent.mkdir(parents=True, exist_ok=True)


@contextmanager
def feedback_connection():
    if USE_POSTGRES:
        # Import lazily so local SQLite installations still work.
        import psycopg
        from psycopg.rows import dict_row

        with psycopg.connect(DATABASE_URL, connect_timeout=10, row_factory=dict_row) as connection:
            yield connection
    else:
        connection = sqlite3.connect(str(FEEDBACK_DB), timeout=15)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA busy_timeout=15000")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()


def feedback_execute(connection, sql: str, params=()):
    # SQL statements in this module use SQLite-style '?' placeholders.
    # psycopg expects '%s' placeholders instead.
    if USE_POSTGRES:
        sql = sql.replace("?", "%s")
    return connection.execute(sql, params)


def initialize_feedback_db():
    with feedback_connection() as connection:
        id_column = "BIGSERIAL PRIMARY KEY" if USE_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS feedback (
                id {id_column},
                name TEXT NOT NULL,
                email TEXT,
                message TEXT NOT NULL,
                created_at TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'new'
            )
        """)
        connection.execute("CREATE INDEX IF NOT EXISTS idx_feedback_created ON feedback(created_at DESC)")
        columns = {
            row["column_name"] if USE_POSTGRES else row["name"]
            for row in (connection.execute(
                "SELECT column_name FROM information_schema.columns WHERE table_schema = current_schema() AND table_name = 'feedback'"
            ).fetchall() if USE_POSTGRES else connection.execute("PRAGMA table_info(feedback)").fetchall())
        }
        for column in ("reply_text", "reply_sent_at", "reply_email_id", "reply_send_state"):
            if column not in columns:
                connection.execute(f"ALTER TABLE feedback ADD COLUMN {column} TEXT")


initialize_feedback_db()


class FeedbackSubmission(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    email: str | None = Field(default=None, max_length=254)
    message: str = Field(min_length=5, max_length=5000)
    website: str | None = Field(default=None, max_length=200)  # honeypot


class FeedbackStatusUpdate(BaseModel):
    status: str = Field(pattern="^(new|read|replied|archived)$")


class FeedbackEmailReply(BaseModel):
    message: str = Field(min_length=1, max_length=5000)


def require_feedback_admin(authorization: str | None):
    if not ADMIN_TOKEN:
        raise HTTPException(status_code=503, detail="Author dashboard is not configured yet.")
    scheme, _, supplied = (authorization or "").partition(" ")
    if scheme.lower() != "bearer" or not secrets.compare_digest(supplied, ADMIN_TOKEN):
        raise HTTPException(status_code=401, detail="Invalid admin access token.")


# ========================================
# CORS
# ========================================

CORS_ORIGINS = [
    origin.strip()
    for origin in os.getenv(
        "CORS_ORIGINS",
        "https://hotpdf.haoyanbing.workers.dev,"
        "http://localhost:5173,"
        "http://localhost:3000,"
        "http://127.0.0.1:5173"
    ).split(",")
    if origin.strip()
]


# Limit concurrent heavy operations.

processing_slots = asyncio.Semaphore(2)


# ========================================
# Temporary file management
# ========================================

def safe_unlink(path: Path):

    try:
        path.unlink(missing_ok=True)

    except OSError:
        logger.warning(
            "Could not remove temporary file",
            exc_info=True
        )


def cleanup_expired_files():

    cutoff = time.time() - TEMP_TTL_SECONDS

    for item in UPLOAD_DIR.iterdir():

        try:

            if not item.is_file():
                continue

            if item.is_symlink():
                continue

            if item.stat().st_mtime < cutoff:
                item.unlink()

        except FileNotFoundError:
            continue

        except OSError:

            logger.warning(
                "Temporary file cleanup failed",
                exc_info=True
            )


async def cleanup_loop():

    while True:

        try:

            await run_in_threadpool(
                cleanup_expired_files
            )

        except Exception:

            logger.exception(
                "Cleanup task failed"
            )

        await asyncio.sleep(
            CLEAN_INTERVAL_SECONDS
        )


@asynccontextmanager
async def lifespan(app: FastAPI):

    cleanup_task = asyncio.create_task(
        cleanup_loop()
    )

    try:
        yield

    finally:

        cleanup_task.cancel()

        with suppress(asyncio.CancelledError):
            await cleanup_task


# ========================================
# FastAPI application
# ========================================

app = FastAPI(lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["*"]
)


# ========================================
# File helpers
# ========================================

def new_filename(prefix: str, extension: str):

    return (
        f"{prefix}_{uuid.uuid4().hex}.{extension}"
    )


def stored_path(
    filename: str,
    pdf_only: bool = False
):

    pattern = (
        PDF_NAME_RE if pdf_only else FILE_NAME_RE
    )

    if not pattern.fullmatch(filename):

        raise HTTPException(
            status_code=400,
            detail="Invalid file identifier."
        )

    path = UPLOAD_DIR / filename

    if not path.is_file():

        raise HTTPException(
            status_code=404,
            detail="File not found or expired."
        )

    try:

        # Refresh expiration time on access.
        os.utime(path, None)

    except FileNotFoundError:

        raise HTTPException(
            status_code=404,
            detail="File not found or expired."
        ) from None

    return path


async def save_pdf_upload(
    upload: UploadFile,
    path: Path
):

    original_name = upload.filename or ""

    if not original_name.lower().endswith(".pdf"):

        raise HTTPException(
            status_code=400,
            detail="Please choose a PDF file."
        )

    total = 0

    try:

        with path.open("xb") as output:

            while True:

                chunk = await upload.read(
                    1024 * 1024
                )

                if not chunk:
                    break

                total += len(chunk)

                if total > MAX_FILE_BYTES:

                    raise HTTPException(
                        status_code=413,
                        detail=(
                            "Each PDF must be no larger "
                            f"than {MAX_FILE_MB} MB."
                        )
                    )

                output.write(chunk)

        if total == 0:

            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty."
            )

        return total

    except Exception:

        safe_unlink(path)
        raise

    finally:

        await upload.close()


def inspect_pdf(path: Path):

    with path.open("rb") as stream:

        header = stream.read(1024)

        if b"%PDF-" not in header:

            raise HTTPException(
                status_code=400,
                detail="This is not a valid PDF file."
            )

    try:

        with fitz.open(str(path)) as pdf:

            if pdf.needs_pass:

                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Password-protected PDFs "
                        "are not supported."
                    )
                )

            count = pdf.page_count

            if count < 1:

                raise HTTPException(
                    status_code=400,
                    detail="The PDF has no pages."
                )

            if count > MAX_PDF_PAGES:

                raise HTTPException(
                    status_code=413,
                    detail=(
                        "A PDF may contain at most "
                        f"{MAX_PDF_PAGES} pages."
                    )
                )

            return count

    except HTTPException:
        raise

    except Exception as exc:

        logger.info(
            "Invalid PDF: %s",
            type(exc).__name__
        )

        raise HTTPException(
            status_code=400,
            detail="Unable to read this PDF file."
        ) from None


# ========================================
# PDF processing helpers
# ========================================

def make_split(
    source: Path,
    output: Path,
    pages: list[int]
):

    with fitz.open(str(source)) as pdf:

        with fitz.open() as result:

            for number in pages:

                if not 1 <= number <= pdf.page_count:

                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"Page {number} is outside "
                            "the PDF page range."
                        )
                    )

                result.insert_pdf(
                    pdf,
                    from_page=number - 1,
                    to_page=number - 1
                )

            result.save(str(output))

    return len(pages)


def make_merge(
    paths: list[Path],
    output: Path
):

    page_count = 0

    with fitz.open() as merged:

        for path in paths:

            with fitz.open(str(path)) as pdf:

                page_count += pdf.page_count

                if page_count > MAX_MERGE_PAGES:

                    raise HTTPException(
                        status_code=413,
                        detail=(
                            "Merged PDF may contain "
                            f"at most {MAX_MERGE_PAGES} pages."
                        )
                    )

                merged.insert_pdf(pdf)

        merged.save(str(output))

    return page_count


def make_word(
    source: Path,
    output: Path
):

    word = Document()

    has_text = False

    with fitz.open(str(source)) as pdf:

        for page in pdf:

            content = page.get_text().strip()

            if content:

                has_text = True

                word.add_paragraph(content)

    if not has_text:

        raise HTTPException(
            status_code=422,
            detail=(
                "This PDF has no extractable text. "
                "Scanned PDFs require OCR."
            )
        )

    word.save(str(output))


def make_compressed(
    source: Path,
    output: Path
):

    command = [
        "gs",
        "-dSAFER",
        "-sDEVICE=pdfwrite",
        "-dCompatibilityLevel=1.4",
        "-dPDFSETTINGS=/ebook",
        "-dDetectDuplicateImages=true",
        "-dNOPAUSE",
        "-dQUIET",
        "-dBATCH",
        f"-sOutputFile={output}",
        str(source)
    ]

    try:

        subprocess.run(
            command,
            check=True,
            capture_output=True,
            timeout=90
        )

    except FileNotFoundError:

        logger.error(
            "Ghostscript executable missing"
        )

        raise HTTPException(
            status_code=503,
            detail=(
                "PDF compression is "
                "temporarily unavailable."
            )
        ) from None

    except subprocess.TimeoutExpired:

        raise HTTPException(
            status_code=504,
            detail="PDF compression timed out."
        ) from None

    except subprocess.CalledProcessError:

        logger.warning(
            "Ghostscript compression failed",
            exc_info=True
        )

        raise HTTPException(
            status_code=422,
            detail="Could not compress this PDF."
        ) from None

    if (
        not output.is_file()
        or output.stat().st_size == 0
    ):

        raise HTTPException(
            status_code=500,
            detail=(
                "Compression did not produce "
                "an output PDF."
            )
        )

    # Avoid returning a larger file.

    if (
        output.stat().st_size
        >= source.stat().st_size
    ):

        shutil.copyfile(
            source,
            output
        )


def make_compressed_to_target(
    source: Path,
    output: Path,
    target_bytes: int
):
    """Try several compression profiles to meet a maximum size.

    Returns True only when the actual output file is within the requested
    byte limit. If no profile reaches the limit, returns the smallest valid
    result found (or the original file if no smaller result is available).
    """

    original_size = source.stat().st_size

    if original_size <= target_bytes:
        shutil.copyfile(source, output)
        return True

    profiles = [
        (150, 75),
        (120, 65),
        (96, 55),
        (72, 40),
    ]

    deadline = time.monotonic() + 100
    best_path = None
    best_size = original_size

    try:
        for index, (resolution, quality) in enumerate(profiles):
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                break

            candidate = output.with_name(
                f"{output.stem}_attempt{index}.pdf"
            )

            command = [
                "gs",
                "-dSAFER",
                "-sDEVICE=pdfwrite",
                "-dCompatibilityLevel=1.4",
                "-dNOPAUSE",
                "-dQUIET",
                "-dBATCH",
                "-dDetectDuplicateImages=true",
                "-dDownsampleColorImages=true",
                "-dDownsampleGrayImages=true",
                "-dDownsampleMonoImages=true",
                "-dColorImageDownsampleType=/Bicubic",
                "-dGrayImageDownsampleType=/Bicubic",
                "-dMonoImageDownsampleType=/Subsample",
                f"-dColorImageResolution={resolution}",
                f"-dGrayImageResolution={resolution}",
                f"-dMonoImageResolution={resolution}",
                "-dAutoFilterColorImages=false",
                "-dAutoFilterGrayImages=false",
                "-dColorImageFilter=/DCTEncode",
                "-dGrayImageFilter=/DCTEncode",
                f"-dJPEGQ={quality}",
                f"-sOutputFile={candidate}",
                str(source),
            ]

            try:
                subprocess.run(
                    command,
                    check=True,
                    capture_output=True,
                    timeout=min(30, remaining),
                )

                if not candidate.is_file() or candidate.stat().st_size == 0:
                    continue

                try:
                    with fitz.open(str(candidate)) as pdf:
                        if pdf.page_count < 1:
                            continue
                except Exception:
                    logger.warning("Invalid PDF from target compression attempt")
                    continue

                candidate_size = candidate.stat().st_size
                if candidate_size < best_size:
                    if best_path is not None:
                        safe_unlink(best_path)
                    best_path = candidate
                    best_size = candidate_size
                    candidate = None

                if best_size <= target_bytes:
                    break

            except subprocess.TimeoutExpired:
                logger.warning("Target compression attempt timed out")
            except subprocess.CalledProcessError:
                logger.warning("Target compression attempt failed")
            finally:
                if candidate is not None:
                    safe_unlink(candidate)

        if best_path is None:
            shutil.copyfile(source, output)
            return False

        shutil.copyfile(best_path, output)
        return output.stat().st_size <= target_bytes

    except FileNotFoundError:
        logger.error("Ghostscript executable missing")
        raise HTTPException(
            status_code=503,
            detail="PDF compression is temporarily unavailable.",
        ) from None
    finally:
        if best_path is not None:
            safe_unlink(best_path)


def make_pdf_images(
    source: Path,
    output: Path,
    image_format: str,
):
    """Render every PDF page into an image and package the pages as a ZIP."""
    image_format = image_format.lower()
    if image_format not in {"jpg", "png"}:
        raise HTTPException(
            status_code=400,
            detail="Image format must be JPG or PNG.",
        )

    image_type = "jpeg" if image_format == "jpg" else "png"

    with fitz.open(str(source)) as pdf:
        if pdf.page_count > MAX_PDF_PAGES:
            raise HTTPException(
                status_code=413,
                detail=f"A PDF may contain at most {MAX_PDF_PAGES} pages.",
            )

        with zipfile.ZipFile(
            output, mode="w", compression=zipfile.ZIP_DEFLATED
        ) as archive:
            for index, page in enumerate(pdf, start=1):
                rect = page.rect
                scale = min(2.0, 2000 / max(rect.width, rect.height, 1))
                pixmap = page.get_pixmap(
                    matrix=fitz.Matrix(scale, scale),
                    alpha=False,
                )
                archive.writestr(
                    f"page-{index:03d}.{image_format}",
                    pixmap.tobytes(image_type),
                )


def make_images_pdf(
    image_contents: list[tuple[str, bytes]],
    output: Path,
):
    """Create a reasonably sized PDF with one page per image.

    Images are downsampled only when their longest side exceeds 2200 pixels,
    then encoded as JPEG to avoid embedding very large original PNG streams.
    Pages use a 150-DPI physical size rather than treating pixels as PDF points.
    """
    max_dimension = 2200
    target_dpi = 150

    with fitz.open() as result:
        for image_name, content in image_contents:
            extension = Path(image_name).suffix.lower()
            image_type = "jpeg" if extension in {".jpg", ".jpeg"} else "png"

            try:
                with fitz.open(stream=content, filetype=image_type) as image:
                    if image.page_count < 1:
                        raise ValueError("empty image")
                    rect = image[0].rect
                    longest_side = max(rect.width, rect.height)
                    scale = min(1.0, max_dimension / max(longest_side, 1))
                    pixmap = image[0].get_pixmap(
                        matrix=fitz.Matrix(scale, scale),
                        alpha=False,
                    )
                    # JPEG is much smaller than many camera/phone PNG images.
                    optimized_image = pixmap.tobytes("jpeg")
                    width = pixmap.width * 72 / target_dpi
                    height = pixmap.height * 72 / target_dpi
            except Exception:
                raise HTTPException(
                    status_code=400,
                    detail=f"Could not read image: {image_name}",
                ) from None

            page = result.new_page(width=width, height=height)
            page.insert_image(page.rect, stream=optimized_image)

        result.save(str(output), garbage=4, deflate=True, clean=True)


def png_preview(
    path: Path,
    page_number: int
):

    with fitz.open(str(path)) as pdf:

        if (
            page_number < 1
            or page_number > pdf.page_count
        ):

            raise HTTPException(
                status_code=400,
                detail="Invalid page number."
            )

        page = pdf[page_number - 1]

        width = max(
            page.rect.width,
            1
        )

        height = max(
            page.rect.height,
            1
        )

        scale = min(
            1.4,
            2000 / max(width, height)
        )

        pixmap = page.get_pixmap(
            matrix=fitz.Matrix(
                scale,
                scale
            ),
            alpha=False
        )

        return pixmap.tobytes("png")


# ========================================
# Office document conversion helpers
# ========================================

def run_libreoffice(input_path: Path, output_dir: Path, target: str):
    """Convert DOCX/PPTX using LibreOffice in headless mode."""
    try:
        subprocess.run(
            ["libreoffice", "-env:UserInstallation=file:///tmp/lo-" + uuid.uuid4().hex,
             "--headless", "--convert-to", target, "--outdir", str(output_dir), str(input_path)],
            check=True, capture_output=True, timeout=120,
        )
    except FileNotFoundError:
        raise HTTPException(status_code=503, detail="Office conversion is unavailable: LibreOffice is not installed.") from None
    except subprocess.TimeoutExpired:
        raise HTTPException(status_code=504, detail="Office conversion timed out.") from None
    except subprocess.CalledProcessError:
        raise HTTPException(status_code=422, detail="Could not convert this Office document.") from None
    output = output_dir / f"{input_path.stem}.{target}"
    if not output.is_file() or output.stat().st_size == 0:
        raise HTTPException(status_code=422, detail="Conversion did not produce an output file.")
    return output


def pdf_to_powerpoint(source: Path, output: Path):
    """Render each PDF page to a slide image. This preserves appearance, not editable text."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    with fitz.open(str(source)) as pdf:
        if pdf.page_count > MAX_PDF_PAGES:
            raise HTTPException(status_code=413, detail=f"A PDF may contain at most {MAX_PDF_PAGES} pages.")
        for page in pdf:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
            image_bytes = pix.tobytes("png")
            from io import BytesIO
            stream = BytesIO(image_bytes)
            slide = prs.slides.add_slide(blank)
            page_ratio = page.rect.width / max(page.rect.height, 1)
            slide_ratio = prs.slide_width / prs.slide_height
            if page_ratio >= slide_ratio:
                width = prs.slide_width
                height = int(width / page_ratio)
                left, top = 0, int((prs.slide_height - height) / 2)
            else:
                height = prs.slide_height
                width = int(height * page_ratio)
                left, top = int((prs.slide_width - width) / 2), 0
            slide.shapes.add_picture(stream, left, top, width=width, height=height)
    prs.save(str(output))


# ========================================
# Health checks
# ========================================

@app.get("/")
def home():

    return {
        "message": "PDF API running"
    }


@app.get("/health")
def health():

    return {
        "status": "ok"
    }


# ========================================
# Upload PDF
# ========================================

@app.post("/upload")
async def upload_pdf(
    file: UploadFile = File(...)
):

    async with processing_slots:

        output_name = new_filename(
            "upload",
            "pdf"
        )

        output = UPLOAD_DIR / output_name

        try:

            await save_pdf_upload(
                file,
                output
            )

            page_count = await run_in_threadpool(
                inspect_pdf,
                output
            )

            return {
                "filename": output_name,
                "pages": page_count,
                "size": round(
                    output.stat().st_size / 1024 / 1024,
                    2
                ),
                "message": "Upload successful"
            }

        except Exception:

            safe_unlink(output)
            raise


# ========================================
# Preview PDF
# ========================================

@app.get("/preview/{filename}/{page}")
async def preview(
    filename: str,
    page: int
):

    source = stored_path(
        filename,
        pdf_only=True
    )

    async with processing_slots:

        image = await run_in_threadpool(
            png_preview,
            source,
            page
        )

    return Response(
        content=image,
        media_type="image/png",
        headers={
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff"
        }
    )


# ========================================
# Split PDF
# ========================================

@app.post("/split")
async def split_pdf(
    filename: str,
    pages: str
):

    source = stored_path(
        filename,
        pdf_only=True
    )

    try:

        numbers = [
            int(item.strip())
            for item in pages.split(",")
        ]

    except ValueError:

        raise HTTPException(
            status_code=400,
            detail=(
                "Pages must be "
                "comma-separated numbers."
            )
        ) from None

    if (
        not numbers
        or len(numbers) > MAX_PDF_PAGES
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid number "
                "of pages selected."
            )
        )

    output_name = new_filename(
        "split",
        "pdf"
    )

    output = UPLOAD_DIR / output_name

    async with processing_slots:

        try:

            page_count = await run_in_threadpool(
                make_split,
                source,
                output,
                numbers
            )

            return {
                "filename": output_name,
                "pages": page_count,
                "size": (
                    f"{output.stat().st_size / 1024:.2f} KB"
                ),
                "message": "Split successful"
            }

        except Exception:

            safe_unlink(output)
            raise


# ========================================
# Merge PDF
# ========================================

@app.post("/merge")
async def merge_pdf(
    files: list[UploadFile] = File(...)
):

    if not 2 <= len(files) <= MAX_MERGE_FILES:

        raise HTTPException(
            status_code=400,
            detail=(
                "Please choose between 2 and "
                f"{MAX_MERGE_FILES} PDFs."
            )
        )

    sources = []

    output_name = new_filename(
        "merged",
        "pdf"
    )

    output = UPLOAD_DIR / output_name

    async with processing_slots:

        try:

            total_bytes = 0

            for file in files:

                temporary = UPLOAD_DIR / new_filename(
                    "upload",
                    "pdf"
                )

                sources.append(temporary)

                size = await save_pdf_upload(
                    file,
                    temporary
                )

                total_bytes += size

                if total_bytes > MAX_MERGE_TOTAL_BYTES:

                    raise HTTPException(
                        status_code=413,
                        detail=(
                            "Total merge size must not "
                            f"exceed {MAX_MERGE_TOTAL_MB} MB."
                        )
                    )

                await run_in_threadpool(
                    inspect_pdf,
                    temporary
                )

            page_count = await run_in_threadpool(
                make_merge,
                sources,
                output
            )

            return {
                "filename": output_name,
                "pages": page_count,
                "size": (
                    f"{output.stat().st_size / 1024:.2f} KB"
                ),
                "message": "Merge successful"
            }

        except Exception:

            safe_unlink(output)
            raise

        finally:

            for path in sources:
                safe_unlink(path)


# ========================================
# Convert PDF to JPG / PNG
# ========================================

@app.post("/pdf_to_images")
async def pdf_to_images(
    filename: str,
    image_format: str = "jpg",
):
    source = stored_path(filename, pdf_only=True)
    image_format = image_format.lower()

    if image_format not in {"jpg", "png"}:
        raise HTTPException(
            status_code=400,
            detail="Image format must be JPG or PNG.",
        )

    output_name = new_filename("converted", "zip")
    output = UPLOAD_DIR / output_name

    async with processing_slots:
        try:
            await run_in_threadpool(
                make_pdf_images, source, output, image_format
            )
            return {
                "filename": output_name,
                "size": f"{output.stat().st_size / 1024:.2f} KB",
                "message": f"PDF pages converted to {image_format.upper()} images.",
            }
        except Exception:
            safe_unlink(output)
            raise


# ========================================
# Convert JPG / PNG images to PDF
# ========================================

@app.post("/images_to_pdf")
async def images_to_pdf(
    files: list[UploadFile] = File(...),
):
    if not 1 <= len(files) <= 20:
        raise HTTPException(
            status_code=400,
            detail="Please choose between 1 and 20 images.",
        )

    allowed_extensions = {".jpg", ".jpeg", ".png"}
    image_contents = []
    total_bytes = 0

    try:
        for file in files:
            original_name = Path(file.filename or "").name
            extension = Path(original_name).suffix.lower()

            if extension not in allowed_extensions:
                raise HTTPException(
                    status_code=400,
                    detail="Only JPG, JPEG, and PNG images are supported.",
                )

            content = await file.read(MAX_FILE_BYTES + 1)
            if not content:
                raise HTTPException(
                    status_code=400,
                    detail=f"The image is empty: {original_name}",
                )
            if len(content) > MAX_FILE_BYTES:
                raise HTTPException(
                    status_code=413,
                    detail=f"Each image must be no larger than {MAX_FILE_MB} MB.",
                )

            total_bytes += len(content)
            if total_bytes > MAX_MERGE_TOTAL_BYTES:
                raise HTTPException(
                    status_code=413,
                    detail=f"Total image size must not exceed {MAX_MERGE_TOTAL_MB} MB.",
                )

            image_contents.append((original_name, content))

        output_name = new_filename("converted", "pdf")
        output = UPLOAD_DIR / output_name

        async with processing_slots:
            try:
                await run_in_threadpool(make_images_pdf, image_contents, output)
                return {
                    "filename": output_name,
                    "size": f"{output.stat().st_size / 1024:.2f} KB",
                    "pages": len(image_contents),
                    "message": "Images converted to PDF successfully.",
                }
            except Exception:
                safe_unlink(output)
                raise
    finally:
        for file in files:
            await file.close()


# ========================================
# Convert PDF to Word
# ========================================

@app.post("/convert_word")
async def convert_word(
    filename: str
):

    source = stored_path(
        filename,
        pdf_only=True
    )

    output_name = new_filename(
        "converted",
        "docx"
    )

    output = UPLOAD_DIR / output_name

    async with processing_slots:

        try:

            await run_in_threadpool(
                make_word,
                source,
                output
            )

            return {
                "filename": output_name,
                "size": (
                    f"{output.stat().st_size / 1024:.2f} KB"
                ),
                "message": "Conversion successful"
            }

        except Exception:

            safe_unlink(output)
            raise


# ========================================
# Edit PDF pages (rotate / delete)
# ========================================

@app.post("/edit_pdf")
async def edit_pdf(
    filename: str,
    rotations: str = "{}",
    delete_pages: str = "",
    page_order: str = "",
    insert_filename: str = "",
    insert_pages: str = "",
    insert_before: int = 1,
):
    """Rotate selected pages and/or delete pages, using 1-based page numbers."""
    source = stored_path(filename, pdf_only=True)

    try:
        rotation_map = json.loads(rotations)
        if not isinstance(rotation_map, dict):
            raise ValueError("rotations must be an object")
        normalized_rotations = {}
        for raw_page, raw_degrees in rotation_map.items():
            page_number = int(raw_page)
            degrees = int(raw_degrees)
            if degrees not in (0, 90, 180, 270):
                raise ValueError("rotation must be 0, 90, 180, or 270")
            normalized_rotations[page_number] = degrees
        deleted = [int(item.strip()) for item in delete_pages.split(",") if item.strip()] if delete_pages.strip() else []
        if len(deleted) != len(set(deleted)):
            raise ValueError("duplicate page numbers")
        requested_order = None
        if page_order.strip():
            requested_order = json.loads(page_order)
            if not isinstance(requested_order, list):
                raise ValueError("page_order must be a list")
            requested_order = [int(item) for item in requested_order]
            if len(requested_order) != len(set(requested_order)):
                raise ValueError("duplicate page order entries")
        inserted_pages = [int(item.strip()) for item in insert_pages.split(",") if item.strip()] if insert_pages.strip() else []
        if len(inserted_pages) != len(set(inserted_pages)):
            raise ValueError("duplicate inserted page numbers")
        if insert_before < 1:
            raise ValueError("insert_before must be positive")
    except (ValueError, TypeError, json.JSONDecodeError):
        raise HTTPException(status_code=400, detail="Invalid page edits. Please refresh and try again.") from None

    insert_source = None
    if inserted_pages:
        if not insert_filename:
            raise HTTPException(status_code=400, detail="Choose a second PDF before inserting pages.")
        insert_source = stored_path(insert_filename, pdf_only=True)
        if insert_source.resolve() == source.resolve():
            raise HTTPException(status_code=400, detail="Choose a different PDF to insert pages from.")

    output_name = new_filename("edited", "pdf")
    output = UPLOAD_DIR / output_name

    async with processing_slots:
        try:
            def apply_edits():
                with fitz.open(str(source)) as pdf:
                    page_count = pdf.page_count
                    valid_pages = set(range(1, page_count + 1))
                    if not set(normalized_rotations).issubset(valid_pages):
                        raise HTTPException(status_code=400, detail="A rotation page number is invalid.")
                    if not set(deleted).issubset(valid_pages):
                        raise HTTPException(status_code=400, detail="A deletion page number is invalid.")
                    if len(deleted) >= page_count:
                        raise HTTPException(status_code=400, detail="A PDF must contain at least one page.")
                    for page_number, degrees in normalized_rotations.items():
                        page = pdf[page_number - 1]
                        page.set_rotation((page.rotation + degrees) % 360)
                    expected_pages = valid_pages - set(deleted)
                    final_order = requested_order if requested_order is not None else sorted(expected_pages)
                    if len(final_order) != len(expected_pages) or set(final_order) != expected_pages:
                        raise HTTPException(
                            status_code=400,
                            detail="Page order must include every kept page exactly once.",
                        )
                    # PyMuPDF select() both removes omitted pages and reorders the kept pages.
                    pdf.select([page_number - 1 for page_number in final_order])
                    if inserted_pages:
                        with fitz.open(str(insert_source)) as extra_pdf:
                            if any(number < 1 or number > extra_pdf.page_count for number in inserted_pages):
                                raise HTTPException(status_code=400, detail="An inserted page number is invalid.")
                            if insert_before > pdf.page_count + 1:
                                raise HTTPException(status_code=400, detail="The insertion position is invalid.")
                            insertion_index = insert_before - 1
                            for page_number in inserted_pages:
                                pdf.insert_pdf(
                                    extra_pdf,
                                    from_page=page_number - 1,
                                    to_page=page_number - 1,
                                    start_at=insertion_index,
                                )
                                insertion_index += 1
                    pdf.save(str(output), garbage=4, deflate=True, clean=True)
                    return pdf.page_count

            page_count = await run_in_threadpool(apply_edits)
            return {
                "filename": output_name,
                "pages": page_count,
                "size": f"{output.stat().st_size / 1024:.2f} KB",
                "message": "PDF edited successfully.",
            }
        except Exception:
            safe_unlink(output)
            raise


# ========================================
# Author feedback endpoints
# ========================================

@app.post("/feedback")
async def submit_feedback(payload: FeedbackSubmission, request: Request):
    # Quietly accept bot submissions into the honeypot without storing them.
    if payload.website and payload.website.strip():
        return {"message": "Thank you for your feedback."}
    name = " ".join(payload.name.split()).strip()
    email = (payload.email or "").strip() or None
    message = payload.message.strip()
    if not name or len(message) < 5:
        raise HTTPException(status_code=422, detail="Please provide your name and a message of at least 5 characters.")
    if email and ("@" not in email or len(email) > 254):
        raise HTTPException(status_code=422, detail="Please enter a valid email address or leave it blank.")
    # Per-process basic abuse guard; use a gateway/WAF for stronger distributed rate limits.
    now = time.time()
    recent = getattr(submit_feedback, "recent", {})
    client_ip = request.client.host if request.client else "unknown"
    recent = {key: value for key, value in recent.items() if now - value < 60}
    if client_ip in recent:
        raise HTTPException(status_code=429, detail="Please wait a minute before sending another message.")
    recent[client_ip] = now
    submit_feedback.recent = recent
    created_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with feedback_connection() as connection:
        feedback_execute(connection,
            "INSERT INTO feedback(name,email,message,created_at,status) VALUES(?,?,?,?, 'new')",
            (name, email, message, created_at),
        )
    return {"message": "Thank you. Your message has been sent to the HotPDF author."}


@app.get("/author/feedback")
def list_feedback(authorization: str | None = Header(default=None), limit: int = 100, offset: int = 0):
    require_feedback_admin(authorization)
    limit = max(1, min(limit, 200))
    offset = max(0, offset)
    with feedback_connection() as connection:
        total = connection.execute("SELECT COUNT(*) AS total FROM feedback").fetchone()["total"]
        rows = feedback_execute(connection,
            "SELECT id,name,email,message,created_at,status,reply_text,reply_sent_at,reply_email_id,reply_send_state FROM feedback ORDER BY id DESC LIMIT ? OFFSET ?",
            (limit, offset),
        ).fetchall()
    return {"total": total, "items": [dict(row) for row in rows]}


@app.patch("/author/feedback/{feedback_id}")
def update_feedback_status(feedback_id: int, payload: FeedbackStatusUpdate, authorization: str | None = Header(default=None)):
    require_feedback_admin(authorization)
    with feedback_connection() as connection:
        cursor = feedback_execute(connection, "UPDATE feedback SET status=? WHERE id=?", (payload.status, feedback_id))
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Feedback message not found.")
    return {"message": "Status updated."}


def send_feedback_email(recipient: str, name: str, message: str):
    """Send a reply through Gmail HTTPS API; return Gmail's message id."""
    if not gmail_api_configured():
        raise HTTPException(status_code=503, detail="Gmail API is not configured.")

    mail = EmailMessage()
    mail["From"] = GMAIL_SENDER
    mail["To"] = recipient
    mail["Subject"] = "A reply to your HotPDF feedback"
    mail["Reply-To"] = GMAIL_SENDER
    mail.set_content(f"Hello {name},\n\nThank you for contacting HotPDF.\n\n{message}\n\nBest regards,\nHotPDF")

    token_data = urlparse.urlencode({
        "client_id": GMAIL_CLIENT_ID,
        "client_secret": GMAIL_CLIENT_SECRET,
        "refresh_token": GMAIL_REFRESH_TOKEN,
        "grant_type": "refresh_token",
    }).encode("utf-8")
    token_request = urlrequest.Request(
        "https://oauth2.googleapis.com/token",
        data=token_data,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    try:
        with urlrequest.urlopen(token_request, timeout=20) as response:
            access_token = json.load(response)["access_token"]
        encoded_mail = base64.urlsafe_b64encode(mail.as_bytes()).decode("ascii")
        api_request = urlrequest.Request(
            "https://gmail.googleapis.com/gmail/v1/users/me/messages/send",
            data=json.dumps({"raw": encoded_mail}).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {access_token}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        with urlrequest.urlopen(api_request, timeout=25) as response:
            result = json.load(response)
        message_id = result.get("id")
        if not message_id:
            raise ValueError("Gmail API response missing message id")
        return message_id
    except (urlerror.HTTPError, urlerror.URLError, TimeoutError, OSError, ValueError, KeyError):
        logger.warning("Gmail API send not confirmed", exc_info=True)
        raise HTTPException(
            status_code=502,
            detail="Gmail API did not confirm sending. Check Gmail Sent folder before retrying.",
        ) from None


@app.post("/author/feedback/{feedback_id}/reply")
async def reply_to_feedback(
    feedback_id: int,
    payload: FeedbackEmailReply,
    authorization: str | None = Header(default=None),
):
    require_feedback_admin(authorization)
    message = payload.message.strip()
    if not message:
        raise HTTPException(status_code=422, detail="Reply cannot be empty.")
    if not gmail_api_configured():
        raise HTTPException(status_code=503, detail="Gmail API is not configured.")

    # Atomic claim prevents two simultaneous requests from sending twice.
    with feedback_connection() as connection:
        cursor = feedback_execute(connection, """
            UPDATE feedback SET reply_send_state='sending'
            WHERE id=? AND reply_sent_at IS NULL
              AND (reply_send_state IS NULL OR reply_send_state='failed')
              AND email IS NOT NULL AND email != ''
        """, (feedback_id,))
        if cursor.rowcount != 1:
            row = feedback_execute(connection, "SELECT id,email,reply_sent_at,reply_send_state FROM feedback WHERE id=?", (feedback_id,)).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="Feedback message not found.")
            if not row["email"]:
                raise HTTPException(status_code=400, detail="This visitor did not provide an email address.")
            raise HTTPException(status_code=409, detail="Reply already sent or sending. Check Gmail Sent folder before retrying.")
        row = feedback_execute(connection, "SELECT name,email FROM feedback WHERE id=?", (feedback_id,)).fetchone()

    try:
        provider_id = await run_in_threadpool(send_feedback_email, row["email"], row["name"], message)
    except HTTPException:
        # Provider errors can be ambiguous; never retry automatically.
        with feedback_connection() as connection:
            feedback_execute(connection, "UPDATE feedback SET reply_send_state='failed' WHERE id=? AND reply_send_state='sending'", (feedback_id,))
        raise

    sent_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with feedback_connection() as connection:
        feedback_execute(connection, """
            UPDATE feedback SET reply_text=?, reply_sent_at=?, reply_email_id=?,
                                reply_send_state='sent', status='replied'
            WHERE id=? AND reply_send_state='sending'
        """, (message, sent_at, provider_id, feedback_id))
    return {"message": "Reply accepted by Gmail API.", "status": "replied",
            "reply_text": message, "reply_sent_at": sent_at, "reply_email_id": provider_id,
            "reply_send_state": "sent"}


@app.delete("/author/feedback/{feedback_id}")
def delete_feedback(feedback_id: int, authorization: str | None = Header(default=None)):
    require_feedback_admin(authorization)
    with feedback_connection() as connection:
        cursor = feedback_execute(connection, "DELETE FROM feedback WHERE id=?", (feedback_id,))
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Feedback message not found.")
    return {"message": "Feedback deleted."}


# ========================================
# Added Office conversion endpoints
# ========================================

@app.post("/convert_word_to_pdf")
async def convert_word_to_pdf(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".docx"):
        raise HTTPException(status_code=400, detail="Please upload a DOCX file.")
    temp_name = new_filename("converted", "docx")
    source = UPLOAD_DIR / temp_name
    output_name = new_filename("converted", "pdf")
    output = UPLOAD_DIR / output_name
    try:
        content = await file.read(MAX_FILE_BYTES + 1)
        if not content or len(content) > MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"File must be no larger than {MAX_FILE_MB} MB.")
        source.write_bytes(content)
        await run_in_threadpool(run_libreoffice, source, UPLOAD_DIR, "pdf")
        generated = UPLOAD_DIR / f"{source.stem}.pdf"
        if generated != output:
            generated.replace(output)
        return {"filename": output_name, "size": f"{output.stat().st_size / 1024:.2f} KB", "message": "Word converted to PDF."}
    finally:
        safe_unlink(source)
        await file.close()


@app.post("/convert_powerpoint_to_pdf")
async def convert_powerpoint_to_pdf(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".pptx"):
        raise HTTPException(status_code=400, detail="Please upload a PPTX file.")
    source = UPLOAD_DIR / new_filename("converted", "pptx")
    output_name = new_filename("converted", "pdf")
    output = UPLOAD_DIR / output_name
    try:
        content = await file.read(MAX_FILE_BYTES + 1)
        if not content or len(content) > MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"File must be no larger than {MAX_FILE_MB} MB.")
        source.write_bytes(content)
        await run_in_threadpool(run_libreoffice, source, UPLOAD_DIR, "pdf")
        generated = UPLOAD_DIR / f"{source.stem}.pdf"
        generated.replace(output)
        return {"filename": output_name, "size": f"{output.stat().st_size / 1024:.2f} KB", "message": "PowerPoint converted to PDF."}
    finally:
        safe_unlink(source)
        await file.close()


@app.post("/convert_powerpoint")
async def convert_powerpoint(filename: str):
    source = stored_path(filename, pdf_only=True)
    output_name = new_filename("converted", "pptx")
    output = UPLOAD_DIR / output_name
    try:
        await run_in_threadpool(pdf_to_powerpoint, source, output)
        return {"filename": output_name, "size": f"{output.stat().st_size / 1024:.2f} KB", "message": "PDF pages exported to PowerPoint slides as images."}
    except Exception:
        safe_unlink(output)
        raise


# ========================================
# Download
# ========================================

@app.get("/download/{filename}")
async def download(
    filename: str
):

    path = stored_path(filename)

    if path.suffix == ".pdf":
        media_type = "application/pdf"
    elif path.suffix == ".zip":
        media_type = "application/zip"
    elif path.suffix == ".pptx":
        media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif path.suffix == ".docx":
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        media_type = "application/octet-stream"

    return FileResponse(
        str(path),
        filename=path.name,
        media_type=media_type,
        headers={
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff"
        }
    )


# ========================================
# Compress PDF
# ========================================

@app.post("/compress")
async def compress_pdf(
    filename: str,
    target_bytes: int | None = None,
):
    source = stored_path(
        filename,
        pdf_only=True
    )

    if target_bytes is not None and target_bytes < 1024:
        raise HTTPException(
            status_code=400,
            detail="Target size must be at least 1 KB.",
        )

    if target_bytes is not None and target_bytes > 100 * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail="Target size must not exceed 100 MB.",
        )

    output_name = new_filename(
        "compressed",
        "pdf"
    )
    output = UPLOAD_DIR / output_name

    async with processing_slots:
        try:
            if target_bytes is None:
                # Preserve the existing standard-compression behavior.
                await run_in_threadpool(
                    make_compressed,
                    source,
                    output,
                )
                target_met = None
                message = "Compression successful"
            else:
                target_met = await run_in_threadpool(
                    make_compressed_to_target,
                    source,
                    output,
                    target_bytes,
                )
                message = (
                    "Target size reached"
                    if target_met
                    else (
                        "Could not reach the requested size. "
                        "The best available result is provided."
                    )
                )

            original_bytes = source.stat().st_size
            compressed_bytes = output.stat().st_size
            original_mb = original_bytes / (1024 * 1024)
            compressed_mb = compressed_bytes / (1024 * 1024)
            reduction = (
                (1 - compressed_bytes / original_bytes) * 100
                if original_bytes
                else 0
            )

            return {
                "filename": output_name,
                "original_size": round(original_mb, 2),
                "compressed_size": round(compressed_mb, 2),
                "original_bytes": original_bytes,
                "compressed_bytes": compressed_bytes,
                "saved": round(max(0, reduction), 1),
                "target_bytes": target_bytes,
                "target_met": target_met,
                "message": message,
            }

        except Exception:
            safe_unlink(output)
            raise
